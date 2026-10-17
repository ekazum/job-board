"""Plain text from an uploaded CV."""

import io
import zipfile
from pathlib import PurePath

import docx
from docx.opc.exceptions import PackageNotFoundError
from pypdf import PdfReader
from pypdf.errors import PdfReadError

SUPPORTED_SUFFIXES = (".pdf", ".docx", ".md", ".txt")


class DocumentText:
    """Extracts the text of a .pdf (pypdf), .docx (python-docx, paragraphs then table cells), .md or .txt file
    (read as UTF-8). ValueError for any other type or a file that cannot be read."""

    @staticmethod
    def extract(filename: str, data: bytes) -> str:
        """The document's text, stripped."""
        suffix = PurePath(filename).suffix.lower()
        if suffix not in SUPPORTED_SUFFIXES:
            raise ValueError(f"unsupported file type {suffix or '(none)'}; use one of {', '.join(SUPPORTED_SUFFIXES)}")
        if suffix == ".pdf":
            return DocumentText._pdf(data)
        if suffix == ".docx":
            return DocumentText._docx(data)
        return data.decode("utf-8", errors="replace").strip()

    @staticmethod
    def _pdf(data: bytes) -> str:
        """Text of every page, in order."""
        try:
            reader = PdfReader(io.BytesIO(data))
            pages = [page.extract_text() or "" for page in reader.pages]
        except (PdfReadError, ValueError, OSError) as error:
            raise ValueError(f"could not read the PDF: {error}") from error
        return "\n\n".join(page.strip() for page in pages if page.strip())

    @staticmethod
    def _docx(data: bytes) -> str:
        """Paragraphs, then each table row as cells separated by " | "."""
        try:
            document = docx.Document(io.BytesIO(data))
        except (zipfile.BadZipFile, PackageNotFoundError, ValueError, KeyError, OSError) as error:
            raise ValueError(f"could not read the Word document: {error}") from error
        lines = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(line for line in lines if line.strip()).strip()
