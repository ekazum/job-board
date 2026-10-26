"""Profile API: upload (.pdf, .docx, .md, .txt; max 5 MB), text extraction, original stored, GET and PATCH."""

import io
from collections.abc import Iterator
from pathlib import Path

import docx
import pytest
from fastapi.testclient import TestClient

from job_board.app_factory import AppFactory
from job_board.config.app_config import AppConfig


def _pdf(text: str) -> bytes:
    """A one-page PDF with this text in Helvetica, offsets computed so pypdf reads it without repair."""
    content = b"BT /F1 12 Tf 72 720 Td (%s) Tj ET" % text.encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n%s\nendstream" % (len(content), content),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n%s\nendobj\n" % (number, body)
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    for offset in offsets:
        out += b"%010d 00000 n \n" % offset
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    return bytes(out)


def _docx(*paragraphs: str) -> bytes:
    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Python"
    table.rows[0].cells[1].text = "10 years"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def client(tmp_path: Path) -> Iterator[TestClient]:
    config = AppConfig(db_path=tmp_path / "data" / "job_board.db", companies=())
    with TestClient(AppFactory(lambda: config, tmp_path / "no-dist").build()) as test_client:
        yield test_client


def _upload(client: TestClient, filename: str, data: bytes) -> dict[str, object]:
    response = client.post("/api/profile", files={"file": (filename, data)})
    assert response.status_code == 200, response.text
    body: dict[str, object] = response.json()
    return body


def test_new_profile_is_blank_with_no_location(client: TestClient) -> None:
    profile = client.get("/api/profile").json()
    assert (profile["cv_text"], profile["filename"], profile["location"], profile["allow_fully_remote"], profile["country"]) == (
        "",
        None,
        "",
        False,
        None,
    )


@pytest.mark.parametrize(("location", "country"), [("Israel", "IL"), (" Tel Aviv-Yafo ", "IL"), ("Berlin", "DE"), ("EMEA", None)])
def test_patch_sets_the_location_and_shows_its_country(client: TestClient, location: str, country: str | None) -> None:
    profile = client.patch("/api/profile", json={"location": location, "allow_fully_remote": True}).json()
    assert (profile["location"], profile["country"], profile["allow_fully_remote"]) == (location.strip(), country, True)
    kept = client.patch("/api/profile", json={"targets": "AI research"}).json()
    assert (kept["location"], kept["allow_fully_remote"], kept["targets"]) == (location.strip(), True, "AI research")


def test_patch_sets_the_job_title_keywords(client: TestClient) -> None:
    include = [" Engineering Director ", "CTO", "cto", "", "AI + engineer"]
    profile = client.patch("/api/profile", json={"include_keywords": include, "exclude_keywords": ["intern"]}).json()
    wanted = ["Engineering Director", "CTO", "AI + engineer"]  # trimmed, without blanks or repeats
    assert (profile["include_keywords"], profile["exclude_keywords"]) == (wanted, ["intern"])
    kept = client.patch("/api/profile", json={"location": "Israel"}).json()
    assert (kept["include_keywords"], kept["exclude_keywords"]) == (wanted, ["intern"])


@pytest.mark.parametrize(
    ("filename", "data", "expected"),
    [
        pytest.param("cv.txt", "Senior ML engineer\nTel Aviv".encode(), "Senior ML engineer\nTel Aviv", id="txt"),
        pytest.param("cv.md", b"# Eli\n\n- LLM infrastructure", "# Eli\n\n- LLM infrastructure", id="md"),
        pytest.param("cv.pdf", _pdf("Senior ML Engineer with LLM experience"), "Senior ML Engineer with LLM experience", id="pdf"),
        pytest.param(
            "CV.DOCX",
            _docx("Senior ML Engineer", "Led LLM serving"),
            "Senior ML Engineer\nLed LLM serving\nPython | 10 years",
            id="docx",
        ),
    ],
)
def test_upload_extracts_text_and_keeps_the_original(client: TestClient, tmp_path: Path, filename: str, data: bytes, expected: str) -> None:
    profile = _upload(client, filename, data)
    assert (profile["cv_text"], profile["filename"], profile["targets"]) == (expected, filename, "")
    assert profile["uploaded_at"] is not None
    assert (tmp_path / "data" / "profile" / filename).read_bytes() == data
    assert client.get("/api/profile").json()["cv_text"] == expected


def test_upload_keeps_targets_and_patch_edits_both(client: TestClient) -> None:
    assert client.patch("/api/profile", json={"targets": "AI research in Israel"}).json()["targets"] == "AI research in Israel"
    uploaded = _upload(client, "cv.txt", b"First CV")
    assert (uploaded["cv_text"], uploaded["targets"]) == ("First CV", "AI research in Israel")
    edited = client.patch("/api/profile", json={"cv_text": "Edited CV"}).json()
    assert (edited["cv_text"], edited["targets"], edited["filename"]) == ("Edited CV", "AI research in Israel", "cv.txt")


def test_path_parts_are_stripped_from_the_stored_name(client: TestClient, tmp_path: Path) -> None:
    profile = _upload(client, "..\\..\\evil name?.txt", b"CV")
    assert profile["filename"] == "evil name_.txt"
    assert (tmp_path / "data" / "profile" / "evil name_.txt").exists()


@pytest.mark.parametrize(
    ("filename", "data", "status", "message"),
    [
        pytest.param("cv.exe", b"MZ", 422, "unsupported file type", id="exe"),
        pytest.param("cv.pdf", b"not a pdf", 422, "could not read the PDF", id="broken-pdf"),
        pytest.param("cv.docx", b"not a zip", 422, "could not read the Word document", id="broken-docx"),
        pytest.param("cv.txt", b"   ", 422, "no text found", id="blank"),
        pytest.param("cv.txt", b"x" * (5 * 1024 * 1024 + 1), 413, "larger than 5 MB", id="over-5MB"),
    ],
)
def test_bad_uploads_are_refused(client: TestClient, filename: str, data: bytes, status: int, message: str) -> None:
    response = client.post("/api/profile", files={"file": (filename, data)})
    assert response.status_code == status
    assert message in response.json()["detail"]
    assert client.get("/api/profile").json()["filename"] is None
